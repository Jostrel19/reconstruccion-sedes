"""Lector del formato APU — presupuesto de obra con la sede como fila divisoria.

Es el formato que usan Aranzazu y Aguadas (comparten plantilla). A diferencia
del de Samaná, aquí NO hay una columna de sede: la sede aparece como una fila
de texto dentro de la misma tabla, y los ítems que vienen debajo le pertenecen
hasta que aparece la siguiente. Hay que leer con estado, no mapeando columnas.

Tres cosas que este formato obliga a resolver, y que un lector ingenuo se come:

1. **No trae código DANE.** El cruce tiene que hacerse por nombre contra el
   catálogo, y eso NUNCA se da por bueno solo: el DANE sale como `propuesto` y
   el registro no se vuelca hasta que alguien lo confirma. La regla del proyecto
   es cruzar por DANE; cuando no hay DANE, lo honesto es decirlo, no inventarlo.

2. **Las hojas pueden venir con las fórmulas rotas.** El archivo de Aguadas trae
   `#N/A` y `#REF!` en el 24 % de las celdas y NINGUNA de sus 14 sedes tiene un
   total calculable. Por eso el lector escoge la hoja con menos errores y reporta
   cada celda rota en vez de leer un cero silencioso.

3. **El total de la sede viene declarado y además se puede sumar.** Manda el
   recalculado (regla del HANDOFF, módulo 1 punto 3); si no coincide con el
   declarado, la diferencia se reporta.
"""
from __future__ import annotations

import json
import re
import unicodedata

import openpyxl

import rutas
from ingesta_esquema import Hallazgo, RegistroIngesta

ERRORES_EXCEL = ("#N/A", "#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#NULL!", "#NUM!")

# Encabezados reales de la plantilla. Se buscan por "contiene" porque vienen con
# puntos, tildes y espacios irregulares ("UND.", "CANT.", "VR. UNITARIO").
ENCABEZADOS = {
    "item": "ITEM",
    "descripcion": "DESCRIPCI",
    "unidad": "UND",
    "cantidad": "CANT",
    "vr_unitario": "UNITARIO",
    "vr_total": "TOTAL",
}

# Filas de cierre: no son ítems ni sedes. Se comparan con límite de palabra,
# nunca como subcadena: «IVA» aparece dentro de «ADHESIVA» y descartaría ítems
# de canaleta adhesiva como si fueran la fila del impuesto.
PIE = ("SUBTOTAL", "COSTO TOTAL", "VALOR TOTAL", "VALOR COSTO DIRECTO",
       "ADMINISTRACION", "IMPREVISTOS?", "UTILIDAD", "IVA")
RE_PIE = re.compile(r"\b(" + "|".join(PIE) + r")\b")

# Divisor de institución educativa. Tolera «INSITUCION», que es como viene
# escrito en el archivo real de Aguadas: un lector estricto pierde la
# institución y, con ella, la única forma de saber a qué I.E. pertenece una
# fila que solo dice «SEDE PRINCIPAL».
# Deliberadamente NO incluye «I E»: en Aranzazu los nombres de sede empiezan
# por «I.E JUAN CRISOSTOMO OSORIO SEDE ...» y se leerían como divisor.
RE_IE = re.compile(r"^(INSTITUCION|INSITUCION|INSTITUICION|CENTRO EDUCATIVO)\b")

# Palabras que no distinguen una sede de otra: aparecen en casi todos los
# nombres y hay que quitarlas antes de comparar. Incluye los dos errores de
# digitación que traen las fuentes reales («EDUCAATIVA», «EEDUCATIVO»).
GENERICAS = {
    "INSTITUCION", "INSTITUCIONES", "EDUCATIVA", "EDUCATIVO", "EDUCAATIVA",
    "EEDUCATIVO", "CENTRO", "ESCUELA", "COLEGIO", "SEDE", "PRINCIPAL",
    "RURAL", "MIXTA", "NUEVA", "CONCENTRACION", "ESCOLAR", "ANEXA", "I", "E", "IE",
}


def _es_error(v) -> bool:
    return isinstance(v, str) and v.strip() in ERRORES_EXCEL


def _norm(s) -> str:
    """Sin tildes, sin puntuación y en mayúsculas — solo para cruzar."""
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"[^A-Z0-9 ]+", " ", s.upper()).strip()


def _distintivo(nombre: str, institucion: str | None = None) -> str:
    """La parte del nombre que de verdad identifica a ESA sede y no a otra.

    Hay que quitar el nombre de la institución, no solo las palabras genéricas.
    El archivo escribe «I.E JUAN CRISOSTOMO OSORIO SEDE CAMELIA BAJA» y el
    catálogo «ESCUELA RURAL CAMELIA BAJA»: ambas deben reducirse a
    «CAMELIA BAJA». Si no se quita la institución, «...OSORIO - SEDE PRINCIPAL»
    del catálogo queda contenida en TODAS las sedes de esa institución y les
    propone el mismo DANE a todas — que es precisamente el error que deja
    registros huérfanos.

    La sede principal se reduce a cadena vacía, y una cadena vacía no cruza
    con nada: es correcto, porque «principal» sin más no identifica nada.
    """
    t = _norm(nombre)
    if institucion:
        for w in _norm(institucion).split():
            if w not in GENERICAS and len(w) > 2:
                t = re.sub(rf"\b{re.escape(w)}\b", " ", t)
    palabras = [w for w in t.split() if w not in GENERICAS]
    return " ".join(palabras).strip()


def _elegir_hoja(wb):
    """La plantilla deja hojas viejas con las fórmulas rotas; se usa la que
    tenga encabezado válido y menos errores."""
    candidatas = []
    for nombre in wb.sheetnames:
        ws = wb[nombre]
        if ws.max_row < 5:
            continue
        fila, mapa = _fila_encabezado(ws)
        if fila is None or "vr_total" not in mapa:
            continue
        errores = sum(
            1
            for r in range(1, min(ws.max_row, 400) + 1)
            for c in range(1, min(ws.max_column, 12) + 1)
            if _es_error(ws.cell(row=r, column=c).value)
        )
        candidatas.append((errores, nombre, fila, mapa))
    if not candidatas:
        return None, None, {}
    errores, nombre, fila, mapa = min(candidatas, key=lambda x: x[0])
    return nombre, fila, mapa


def _fila_encabezado(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        textos = {
            c: _norm(ws.cell(row=r, column=c).value)
            for c in range(1, min(ws.max_column, 14) + 1)
        }
        if not any("DESCRIPCI" in t for t in textos.values()):
            continue
        mapa = {}
        for campo, buscado in ENCABEZADOS.items():
            for c, t in textos.items():
                if buscado in t and c not in mapa.values():
                    mapa[campo] = c
                    break
        # "VR. TOTAL" y "COSTO TOTAL" se parecen: el total siempre va a la derecha
        # del valor unitario, así que se corrige si quedó mal asignado.
        if "vr_unitario" in mapa and "vr_total" in mapa and mapa["vr_total"] < mapa["vr_unitario"]:
            mapa["vr_total"] = mapa["vr_unitario"] + 1
        return r, mapa
    return None, {}


def _texto_fila(ws, r, mapa):
    """La descripción migra entre columnas por las celdas combinadas de la
    plantilla: se busca en la columna del encabezado y en la anterior."""
    cols = [mapa.get("descripcion")]
    if mapa.get("descripcion"):
        cols.append(mapa["descripcion"] - 1)
    for c in cols:
        if not c:
            continue
        v = ws.cell(row=r, column=c).value
        if isinstance(v, str) and v.strip() and not _es_error(v):
            return v.strip()
    return ""


def _catalogo_municipio(municipio: str):
    cat = json.loads(rutas.CATALOGO_JSON.read_text(encoding="utf-8"))
    obj = _norm(municipio)
    return [s for s in cat["sedes"] if _norm(s.get("municipio")) == obj]


def _proponer_dane(nombre_sede: str, institucion: str | None, sedes_cat: list):
    """Devuelve (dane, nombre_catalogo, motivo) o (None, None, motivo).

    Nunca devuelve «la más parecida»: o hay una única candidata defendible, o
    se reporta como hallazgo para que lo resuelva una persona.
    """
    clave = _distintivo(nombre_sede, institucion)

    if not clave:
        # Es una «SEDE PRINCIPAL»: por sí sola no identifica nada, pero con la
        # institución sí. En el archivo de Aguadas hay tres filas que dicen solo
        # «SEDE PRINCIPAL» y son de tres instituciones distintas; dos de ellas
        # están en el lote priorizado. Sin este cruce se perderían las dos.
        if not institucion:
            return None, None, "dice solo «sede principal» y la fila no trae institución que la identifique"
        # Se compara también sin espacios: el catálogo escribe «RIOARRIBA» en una
        # palabra y el archivo «RIO ARRIBA». Es la misma institución.
        clave_ie = _distintivo(institucion)
        plano = clave_ie.replace(" ", "")
        principales = [
            s for s in sedes_cat
            if not _distintivo(s.get("sede"), s.get("institucion"))
            and _distintivo(s.get("institucion")).replace(" ", "") == plano
        ]
        if len(principales) == 1:
            return (int(principales[0]["dane_sede"]), principales[0]["sede"],
                    f"sede principal de «{institucion.strip()}»")
        if len(principales) > 1:
            cands = " / ".join(f'{s["dane_sede"]} «{s["sede"]}»' for s in principales)
            return None, None, (f"«{institucion.strip()}» tiene {len(principales)} sedes que el "
                                f"catálogo deja como principal: {cands}. Hay que elegir a mano")
        return None, None, (f"dice solo «sede principal» y «{clave_ie}» no corresponde a ninguna "
                            f"institución del catálogo en este municipio")

    exactos = [s for s in sedes_cat
               if _distintivo(s.get("sede"), s.get("institucion")) == clave]
    if len(exactos) == 1:
        return int(exactos[0]["dane_sede"]), exactos[0]["sede"], "coincidencia exacta del nombre distintivo"
    if len(exactos) > 1:
        return None, None, f"{len(exactos)} sedes del catálogo comparten ese nombre"

    # Sin coincidencia exacta solo se acepta contención de palabras completas,
    # y únicamente si el nombre distintivo es largo: «SAN RAFAEL» no puede
    # cruzarse con «SAN ANTONIO» por compartir una palabra.
    parciales = []
    for s in sedes_cat:
        otro = _distintivo(s.get("sede"), s.get("institucion"))
        if not otro or len(clave) < 6 or len(otro) < 6:
            continue
        if re.search(rf"\b{re.escape(clave)}\b", otro) or re.search(rf"\b{re.escape(otro)}\b", clave):
            parciales.append(s)
    if len(parciales) == 1:
        return int(parciales[0]["dane_sede"]), parciales[0]["sede"], "coincidencia parcial del nombre distintivo"
    if len(parciales) > 1:
        return None, None, f"{len(parciales)} candidatas por coincidencia parcial"
    return None, None, f"«{clave}» no corresponde a ninguna sede del catálogo en este municipio"


def leer(ruta: str, municipio: str):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    registros: list[RegistroIngesta] = []
    hallazgos: list[Hallazgo] = []

    hoja, fila_enc, mapa = _elegir_hoja(wb)
    if not hoja:
        hallazgos.append(Hallazgo("error", "Ninguna hoja tiene un encabezado de presupuesto reconocible", ruta))
        return registros, hallazgos

    ws = wb[hoja]
    descartadas = [n for n in wb.sheetnames if n != hoja and ws.max_row > 5]
    if len(wb.sheetnames) > 1:
        hallazgos.append(Hallazgo(
            "advertencia",
            f"Se leyó la hoja «{hoja}». Las demás hojas del libro no se procesaron — "
            f"verificar que no contengan presupuesto adicional",
            hoja,
        ))

    col = lambda r, campo: ws.cell(row=r, column=mapa[campo]).value if campo in mapa else None

    sedes: list[dict] = []
    actual = None
    institucion = None

    for r in range(fila_enc + 1, ws.max_row + 1):
        item = col(r, "item")
        unidad, cantidad = col(r, "unidad"), col(r, "cantidad")
        vr_unit, vr_total = col(r, "vr_unitario"), col(r, "vr_total")
        texto = _texto_fila(ws, r, mapa)
        tnorm = _norm(texto)
        inorm = _norm(item)

        # filas de cierre del presupuesto (A, U, subtotales).
        # El instrumento ALERTA sobre A y U fuera de umbral, no bloquea: no hay
        # tope legal de AIU, así que la severidad nunca es 'error'.
        if RE_PIE.search(tnorm) or RE_PIE.search(inorm):
            if "ADMINISTRACION" in inorm or "ADMINISTRACION" in tnorm:
                if isinstance(vr_unit, (int, float)) and vr_unit > 0.12:
                    hallazgos.append(Hallazgo(
                        "advertencia",
                        f"Administración al {vr_unit:.0%} del costo directo, por encima del umbral "
                        f"de alerta (12 %). La referencia de mercado sin imprevistos es 10 %. "
                        f"No hay tope legal, pero el municipio debe sustentarlo",
                        f"{hoja}!fila{r}",
                    ))
            if "UTILIDAD" in inorm or "UTILIDAD" in tnorm:
                if isinstance(vr_unit, (int, float)) and vr_unit > 0.08:
                    hallazgos.append(Hallazgo(
                        "advertencia",
                        f"Utilidad al {vr_unit:.0%}, por encima del umbral de alerta (8 %)",
                        f"{hoja}!fila{r}",
                    ))
            continue

        # Divisor de institución educativa. Puede venir en la columna de
        # descripción o en la de ITEM: la plantilla lo pone en una u otra según
        # la hoja, y si se busca solo en una se pierde la institución entera.
        # Un divisor de institución nunca lleva número de orden: eso distingue
        # «INSTITUCIÓN EDUCATIVA RIO ARRIBA» (divisor) de una sede numerada.
        texto_item = item.strip() if isinstance(item, str) else ""
        candidato_ie = texto or texto_item
        if (candidato_ie and not isinstance(item, (int, float))
                and unidad in (None, "") and cantidad in (None, "")
                and RE_IE.match(_norm(candidato_ie))):
            institucion = candidato_ie
            continue

        # divisor de sede: número de orden en ITEM y sin unidad/cantidad
        es_divisor = (
            isinstance(item, (int, float))
            and not isinstance(item, bool)
            and float(item) == int(item)
            and unidad in (None, "")
            and cantidad in (None, "")
            and texto
        )
        if es_divisor:
            actual = {
                "sede": texto, "institucion": institucion, "fila": r,
                "items_ok": 0, "items_rotos": 0, "suma": 0.0,
                "declarado": vr_total if isinstance(vr_total, (int, float)) else None,
            }
            sedes.append(actual)
            continue

        # fila de ítem: tiene unidad o cantidad
        if actual is None or (unidad in (None, "") and cantidad in (None, "")):
            continue

        if any(_es_error(v) for v in (vr_total, vr_unit, cantidad, texto)):
            actual["items_rotos"] += 1
            hallazgos.append(Hallazgo(
                "error",
                f"Fórmula rota en el ítem ({texto or 'sin descripción'}): la celda trae un error de Excel",
                f"{hoja}!fila{r}",
            ))
            continue

        if isinstance(vr_total, (int, float)):
            actual["items_ok"] += 1
            actual["suma"] += float(vr_total)

    if not sedes:
        hallazgos.append(Hallazgo("error", f"No se reconoció ninguna sede en la hoja «{hoja}»", hoja))
        return registros, hallazgos

    sedes_cat = _catalogo_municipio(municipio)

    for s in sedes:
        ref = f"{hoja}!fila{s['fila']}"

        if s["items_rotos"] and s["items_ok"] == 0:
            hallazgos.append(Hallazgo(
                "error",
                f"«{s['sede']}» no tiene ni un solo ítem legible ({s['items_rotos']} con fórmula rota): "
                f"no se puede establecer su costo",
                ref,
            ))
            continue

        if s["items_rotos"]:
            hallazgos.append(Hallazgo(
                "advertencia",
                f"«{s['sede']}» tiene {s['items_rotos']} ítems con fórmula rota que quedaron fuera "
                f"de la suma. El costo leído está incompleto",
                ref,
            ))

        # manda el recalculado, no el declarado
        if isinstance(s["declarado"], (int, float)) and s["declarado"] > 0:
            dif = s["suma"] - s["declarado"]
            if abs(dif) > max(1.0, s["declarado"] * 0.005):
                hallazgos.append(Hallazgo(
                    "advertencia",
                    f"«{s['sede']}»: el total declarado es ${s['declarado']:,.0f} y la suma de sus "
                    f"ítems da ${s['suma']:,.0f} (diferencia ${dif:,.0f}). Manda el recalculado",
                    ref,
                ))

        dane, nombre_cat, motivo = _proponer_dane(s["sede"], s["institucion"], sedes_cat)
        if dane:
            hallazgos.append(Hallazgo(
                "advertencia",
                f"«{s['sede']}» no trae DANE en el archivo. Se propone {dane} "
                f"(«{nombre_cat}») por {motivo}. NO se vuelca hasta confirmarlo",
                ref,
            ))
        else:
            hallazgos.append(Hallazgo(
                "error",
                f"«{s['sede']}» no trae DANE y no se pudo proponer uno: {motivo}",
                ref,
            ))

        registros.append(RegistroIngesta(
            dane_sede=None,                 # sin DANE en el archivo: nada entra en firme
            municipio=municipio,
            costo_directo=s["suma"],
            sede=s["sede"],
            archivo_origen=ruta,
            dane_origen="propuesto",
            dane_propuesto=dane,
            detalle={
                "hoja": hoja, "fila": s["fila"], "institucion": s["institucion"],
                "items_ok": s["items_ok"], "items_rotos": s["items_rotos"],
                "total_declarado": s["declarado"], "nombre_catalogo": nombre_cat,
            },
        ))

    return registros, hallazgos

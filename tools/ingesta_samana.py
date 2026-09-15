"""Lector del formato de SAMANÁ.

No es una tabla de ítems (cantidad x valor unitario): es un diagnóstico por
sede, una hoja por institución educativa. Cada fila de sede ya trae un valor
estimado de materiales y uno de mano de obra; el costo directo es la suma de
los dos.

El encabezado NO está en la misma columna en todas las hojas del mismo libro
-- la hoja "I.E BERLIN" viene corrida una columna a la derecha respecto a las
demás 7. Por eso el encabezado se busca por texto ("CODIGO DANE"), nunca por
letra de columna fija (regla del proyecto: nada de rangos ni columnas fijas).
"""
from __future__ import annotations

import openpyxl

from ingesta_esquema import Hallazgo, RegistroIngesta

# Fragmento de texto que identifica cada columna, tal como aparece en el
# encabezado real. Se busca "contiene", no "es igual a": los encabezados
# reales traen espacios y mayúsculas irregulares.
ENCABEZADOS = {
    "sede": "SEDE",
    "dane": "CODIGO DANE",
    "materiales": "VALOR ESTIMADO DE LA LISTA DE LOS MATERIALES",
    "mano_obra": "VALOR ESTIMADO MANO DE OBRA",
}


def _fila_encabezado(ws):
    """Busca en las primeras filas la que trae 'CODIGO DANE' y arma columna por campo."""
    for r in range(1, min(ws.max_row, 10) + 1):
        textos = {
            c: str(ws.cell(row=r, column=c).value or "").strip().upper()
            for c in range(1, ws.max_column + 1)
        }
        if any("CODIGO DANE" in t for t in textos.values()):
            mapa = {}
            for campo, buscado in ENCABEZADOS.items():
                for c, t in textos.items():
                    if buscado in t:
                        mapa[campo] = c
                        break
            return r, mapa
    return None, {}


def leer(ruta: str, municipio: str = "SAMANÁ"):
    wb = openpyxl.load_workbook(ruta, data_only=True)
    registros: list[RegistroIngesta] = []
    hallazgos: list[Hallazgo] = []

    for nombre_hoja in wb.sheetnames:
        ws = wb[nombre_hoja]
        fila_enc, mapa = _fila_encabezado(ws)
        if fila_enc is None or "dane" not in mapa:
            hallazgos.append(Hallazgo("error", 'No se encontró un encabezado "CODIGO DANE" en la hoja', nombre_hoja))
            continue

        vacias_seguidas = 0
        for r in range(fila_enc + 1, ws.max_row + 1):
            dane = ws.cell(row=r, column=mapa["dane"]).value
            if dane in (None, ""):
                vacias_seguidas += 1
                # Al menos un par de hojas traen una segunda tabla apilada más
                # abajo (p. ej. "I.E PIO XII": una lista de precios de
                # materiales después de un salto en blanco, sin encabezado
                # propio). Dos filas vacías seguidas cierran la tabla de sedes;
                # lo que venga después no se procesa como si fueran sedes.
                if vacias_seguidas >= 2:
                    break
                continue
            vacias_seguidas = 0

            try:
                dane_int = int(dane)
            except (TypeError, ValueError):
                hallazgos.append(Hallazgo("error", f"DANE no numérico: {dane!r}", f"{nombre_hoja}!fila{r}"))
                continue

            if len(str(dane_int)) != 12:
                hallazgos.append(Hallazgo(
                    "advertencia",
                    f"DANE con {len(str(dane_int))} dígitos, se esperaban 12 — revisar antes de cruzar contra el catálogo",
                    f"{nombre_hoja}!fila{r}",
                ))

            def valor(campo):
                col = mapa.get(campo)
                if not col:
                    return 0
                v = ws.cell(row=r, column=col).value
                return v if isinstance(v, (int, float)) else 0

            materiales = valor("materiales")
            mano_obra = valor("mano_obra")
            costo_directo = materiales + mano_obra

            sede_col = mapa.get("sede")
            sede = str(ws.cell(row=r, column=sede_col).value or "").strip() if sede_col else nombre_hoja

            if costo_directo == 0:
                hallazgos.append(Hallazgo(
                    "advertencia",
                    "Sin valor estimado de materiales ni mano de obra — probablemente sede sin afectación, confirmar antes de reportarla como sin costo",
                    f"{nombre_hoja}!fila{r}",
                ))

            registros.append(RegistroIngesta(
                dane_sede=dane_int,
                municipio=municipio,
                costo_directo=costo_directo,
                sede=sede,
                archivo_origen=ruta,
                detalle={"hoja": nombre_hoja, "fila": r, "materiales": materiales, "mano_obra": mano_obra},
            ))

    return registros, hallazgos

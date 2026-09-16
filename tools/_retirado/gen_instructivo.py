"""Genera el instructivo para las alcaldías en formato Word."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import rutas

AZUL = RGBColor(0x12, 0x39, 0x5F)
GRIS = RGBColor(0x5B, 0x68, 0x74)
SALIDA = rutas.INSTRUCTIVO

doc = Document()

for s in doc.sections:
    s.top_margin = Cm(1.8)
    s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(2.0)
    s.right_margin = Cm(2.0)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08


def p(texto="", size=10, bold=False, color=None, align=None, space_before=0, space_after=5):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.space_after = Pt(space_after)
    if align:
        par.alignment = align
    if texto:
        r = par.add_run(texto)
        r.font.size = Pt(size)
        r.bold = bold
        if color:
            r.font.color.rgb = color
    return par


def mixto(partes, size=10, space_before=0, space_after=5, align=None):
    """partes = [(texto, negrita), ...]"""
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(space_before)
    par.paragraph_format.space_after = Pt(space_after)
    if align:
        par.alignment = align
    for texto, bold in partes:
        r = par.add_run(texto)
        r.font.size = Pt(size)
        r.bold = bold
    return par


def titulo(n, texto):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(10)
    par.paragraph_format.space_after = Pt(4)
    r = par.add_run(f"{n}. {texto}")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = AZUL
    return par


def vineta(texto_partes, size=10):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.space_after = Pt(3)
    par.paragraph_format.left_indent = Cm(0.7)
    for texto, bold in texto_partes:
        r = par.add_run(texto)
        r.font.size = Pt(size)
        r.bold = bold
    return par


def sombrear(celda, hex_color):
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tc.append(shd)


# ------------------------------------------------------------------ encabezado
p("GOBERNACIÓN DE CALDAS — SECRETARÍA DE EDUCACIÓN", 9, True, AZUL,
  WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
p("Dirección de Planeación — Infraestructura Educativa", 9, False, GRIS,
  WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
p("INSTRUCTIVO PARA LAS ALCALDÍAS", 15, True, None, WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
p("Radicación del informe técnico-presupuestal por sede educativa", 10.5, False, GRIS,
  WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
p("Reparación de infraestructura educativa afectada por el sismo del 10 de agosto de 2026",
  9, False, GRIS, WD_ALIGN_PARAGRAPH.CENTER, space_after=10)

# ------------------------------------------------------------------- qué y por qué
titulo(1, "Qué se le solicita")
mixto([
    ("La Secretaría de Educación estimó de manera preliminar el costo de reparación de las sedes "
     "educativas oficiales afectadas. Esa estimación es un orden de magnitud: ", False),
    ("no constituye presupuesto de obra", True),
    (". Con este instrumento el municipio presenta el presupuesto real, sede por sede, con "
     "cantidades y valores unitarios.", False),
])
mixto([
    ("Debe pronunciarse sobre ", False),
    ("todas", True),
    (" las sedes oficiales de su municipio: indicar cuáles presentan afectación que requiera "
     "inversión y cuáles no. Solo las que presenten afectación llevan presupuesto.", False),
])

# --------------------------------------------------------------------- acceso
titulo(2, "Cómo ingresar")
vineta([("Use el enlace que recibió por correo institucional. ", False),
        ("Es propio de su municipio y ya trae cargadas sus sedes.", True)])
vineta([("Ábralo completo. Si el correo lo parte en dos líneas, cópielo entero antes de pegarlo "
         "en el navegador.", False)])
vineta([("Trabaje siempre desde ", False), ("el mismo computador y el mismo navegador", True),
        (": lo que va diligenciando queda guardado allí. No use ventanas de incógnito.", False)])
vineta([("No se requiere usuario ni contraseña.", False)])

# ------------------------------------------------------------------ paso a paso
titulo(3, "Paso a paso")

mixto([("Paso 1. Revise la lista de sedes.", True)], space_before=3, space_after=2)
p("Al entrar verá todas las sedes de su municipio con la clasificación de la visita técnica del "
  "27 de agosto de 2026. Esa clasificación es informativa y no se modifica.", space_after=5)

mixto([("Paso 2. Confirme las sedes sin afectación.", True)], space_before=3, space_after=2)
mixto([
    ("Las sedes que el censo clasificó como ", False), ("tipo 5 · sin afectación", True),
    (" vienen pre-marcadas. Puede confirmarlas todas de una vez con el botón que aparece arriba "
     "y concentrarse en las que sí requieren presupuesto. Si alguna de ellas sí presenta "
     "afectación, ábrala y corrija la declaración.", False),
])

mixto([("Paso 3. Diligencie las sedes afectadas.", True)], space_before=3, space_after=2)
vineta([("Describa la afectación y adjunte al menos una fotografía.", False)])
vineta([("Suba una ", False), ("imagen de su firma", True),
        (" —una foto o escaneo de la firma sobre papel blanco—. Queda guardada y se aplica a "
         "todas las sedes del municipio: no hay que subirla cada vez. Aparecerá en el documento "
         "final junto a la firma de la Secretaría.", False)])
vineta([("En el presupuesto, ", False), ("cada fila es un trabajo que se va a ejecutar", True),
        (", no el daño que se observa. El daño ya lo describió arriba.", False)])
vineta([("Si el trabajo se puede medir: indique la unidad, cuántas van y cuánto cuesta cada una. "
         "Por ejemplo, «Reposición de cubierta en teja», unidad m², cantidad 120, valor unitario "
         "$85.000.", False)])
vineta([("Si no se puede desglosar: use la unidad ", False), ("gl (global)", True),
        (", cantidad ", False), ("1", True),
        (", y escriba el costo completo del trabajo en «valor unitario». Por ejemplo, "
         "«Reconstrucción de cocina y restaurante escolar», gl, 1, $80.000.000.", False)])
vineta([("El valor total de cada fila y el costo directo ", False),
        ("los calcula el sistema", True), (". No los escriba.", False)])
vineta([("Puede escribir las cifras con puntos, como se acostumbra: ", False),
        ("1.500.000", True), (".", False)])
vineta([("Sobre el costo directo se agregan dos porcentajes que el municipio define: ", False),
        ("Administración", True),
        (" —lo que cuesta administrar la obra: personal, pólizas, transporte, servicios e "
         "impuestos— y ", False),
        ("Utilidad", True), (" —la ganancia del contratista que ejecute—.", False)])
vineta([("La referencia de mercado es ", False), ("10 % de administración y 5 % de utilidad", True),
        (". No existe un tope legal en Colombia: el municipio los define y debe poder "
         "sustentarlos. No se manejan imprevistos y el IVA del 19 % sobre la utilidad lo "
         "calcula el sistema.", False)])
vineta([("El plazo son los días que toma ", False), ("ejecutar la obra una vez contratada", True),
        (". No incluya el tiempo del proceso de contratación.", False)])
vineta([("Puede guardar borrador y continuar después.", False)])

mixto([("Paso 4. Envíe a la Secretaría.", True)], space_before=3, space_after=2)
vineta([("Mientras diligencia ", False), ("no se descarga nada", True),
        (". Puede trabajar durante varios días: todo queda guardado.", False)])
vineta([("Cuando termine, oprima ", False), ("«Enviar a la Secretaría»", True),
        (" en el tablero. Se descargan los archivos de todas las sedes de una vez y se abre la "
         "carpeta donde debe soltarlos.", False)])
vineta([("El navegador le pedirá permiso para descargar varios archivos: ", False),
        ("acepte", True), (".", False)])
vineta([("Arrastre todos los archivos descargados a la carpeta que se abrió. La primera vez le "
         "pedirá verificar su correo con un código que llega al mismo correo de la alcaldía.",
         False)])
vineta([("Si diligencia más sedes después, vuelva a oprimir «Enviar»: solo se descargan las "
         "nuevas.", False)])

# ------------------------------------------------------------------- fechas
titulo(4, "Plazo")
tabla = doc.add_table(rows=2, cols=2)
tabla.style = "Table Grid"
tabla.alignment = WD_TABLE_ALIGNMENT.LEFT
datos = [("Cierre de la recolección", "Viernes 11 de septiembre de 2026"),
         ("Modalidad", "Ejercicio de prueba para validar el instrumento")]
for i, (a, b) in enumerate(datos):
    tabla.cell(i, 0).text = a
    tabla.cell(i, 1).text = b
    sombrear(tabla.cell(i, 0), "EEF2F6")
    for j in range(2):
        for par in tabla.cell(i, j).paragraphs:
            par.paragraph_format.space_after = Pt(1)
            for r in par.runs:
                r.font.size = Pt(9.5)
                r.bold = (j == 0)
tabla.columns[0].width = Cm(5.2)
tabla.columns[1].width = Cm(11.0)

# --------------------------------------------------------------- advertencias
titulo(5, "Tenga en cuenta")
vineta([("Si su declaración difiere de la del censo —por ejemplo, si reporta afectación en una "
         "sede clasificada sin afectación— ", False),
        ("el sistema le pedirá justificarlo", True),
        (". No es un impedimento: la Secretaría revisará el caso.", False)])
vineta([("El código DANE no se digita: se hereda del catálogo oficial. Así se evita que un "
         "registro quede sin poder cruzarse.", False)])
vineta([("Si una sede no aparece en la lista, ", False), ("no la omita", True),
        (": informe a la Secretaría para verificarla contra la base maestra.", False)])
vineta([("Puede generar el formato en PDF de cada sede radicada desde el mismo aplicativo.", False)])
vineta([("El presupuesto será verificado por la Dirección de Planeación, que podrá aprobarlo, "
         "aprobarlo parcialmente o devolverlo con observaciones.", False)])

# ------------------------------------------------------------------- contacto
titulo(6, "Soporte")
p("Dirección de Planeación — Infraestructura Educativa, Secretaría de Educación de Caldas. "
  "Si extravió el enlace de su municipio, solicítelo por el mismo correo institucional de la "
  "alcaldía; no se atienden solicitudes desde correos personales.", space_after=10)

p("Secretaría de Educación de Caldas — Dirección de Planeación", 8, False, GRIS,
  WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=0)

rutas.asegurar_directorios()
doc.save(SALIDA)
print("Escrito:", SALIDA)
print("Párrafos:", len(doc.paragraphs), "| Tablas:", len(doc.tables))

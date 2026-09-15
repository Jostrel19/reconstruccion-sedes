"""Resumen ejecutivo de una página para la aprobación del instrumento."""
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import rutas

AZUL = RGBColor(0x12, 0x39, 0x5F)
GRIS = RGBColor(0x5B, 0x68, 0x74)
ROJO = RGBColor(0xA3, 0x2B, 0x28)
SALIDA = rutas.DOCS / "RESUMEN EJECUTIVO - Instrumento presupuestal.docx"

doc = Document()
for s in doc.sections:
    s.top_margin = Cm(1.6); s.bottom_margin = Cm(1.4)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)

n = doc.styles["Normal"]
n.font.name = "Calibri"; n.font.size = Pt(10)
n.paragraph_format.space_after = Pt(5); n.paragraph_format.line_spacing = 1.06


def p(txt="", size=10, bold=False, color=None, align=None, sb=0, sa=5):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(sb); par.paragraph_format.space_after = Pt(sa)
    if align: par.alignment = align
    if txt:
        r = par.add_run(txt); r.font.size = Pt(size); r.bold = bold
        if color: r.font.color.rgb = color
    return par


def mixto(partes, sb=0, sa=5, size=10):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(sb); par.paragraph_format.space_after = Pt(sa)
    for t, b in partes:
        r = par.add_run(t); r.font.size = Pt(size); r.bold = b
    return par


def titulo(txt):
    par = doc.add_paragraph()
    par.paragraph_format.space_before = Pt(9); par.paragraph_format.space_after = Pt(3)
    r = par.add_run(txt); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = AZUL
    return par


def vineta(partes, size=9.5):
    par = doc.add_paragraph(style="List Bullet")
    par.paragraph_format.space_after = Pt(2); par.paragraph_format.left_indent = Cm(0.65)
    for t, b in partes:
        r = par.add_run(t); r.font.size = Pt(size); r.bold = b
    return par


def sombrear(celda, color):
    tc = celda._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), color)
    tc.append(shd)


def tabla(datos, anchos, encabezado=True):
    t = doc.add_table(rows=len(datos), cols=len(datos[0]))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, fila in enumerate(datos):
        for j, val in enumerate(fila):
            c = t.cell(i, j); c.text = str(val)
            if encabezado and i == 0: sombrear(c, "EEF2F6")
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(1)
                for r in par.runs:
                    r.font.size = Pt(9)
                    r.bold = (encabezado and i == 0)
    for j, a in enumerate(anchos):
        t.columns[j].width = Cm(a)
    return t


# ---------------------------------------------------------------- encabezado
p("GOBERNACIÓN DE CALDAS — SECRETARÍA DE EDUCACIÓN", 9, True, AZUL,
  WD_ALIGN_PARAGRAPH.CENTER, sa=0)
p("Dirección de Planeación — Infraestructura Educativa", 9, False, GRIS,
  WD_ALIGN_PARAGRAPH.CENTER, sa=8)
p("RESUMEN EJECUTIVO", 14, True, None, WD_ALIGN_PARAGRAPH.CENTER, sa=1)
p("Instrumento de recolección del informe técnico-presupuestal por sede educativa",
  10, False, GRIS, WD_ALIGN_PARAGRAPH.CENTER, sa=1)
p("Sismo del 10 de agosto de 2026 · Manizales, 9 de septiembre de 2026",
  9, False, GRIS, WD_ALIGN_PARAGRAPH.CENTER, sa=9)

# ---------------------------------------------------------------- 1
titulo("1. Qué se resuelve")
mixto([
    ("La estimación de costos del 8 de septiembre valoró en ", False),
    ("$23.348.062.500", True),
    (" la reparación de 456 sedes. El oficio que la acompaña advierte que es una estimación "
     "paramétrica y que ", False),
    ("no constituye presupuesto de obra", True),
    (". Para sustentar la cifra ante el Ministerio, la UNGRD o Hacienda hace falta el presupuesto "
     "real de cada municipio, sede por sede: 975 sedes en 26 municipios.", False),
])
mixto([
    ("El instrumento sustituye el envío de formatos en Word por correo. Entrega la información ",
     False),
    ("ya estructurada y comparable contra la estimación de la Secretaría", True),
    (", sin digitación intermedia.", False),
])

# ---------------------------------------------------------------- 2
titulo("2. Cómo funciona")
vineta([("Cada alcaldía recibe un enlace propio con sus sedes ya cargadas desde la base maestra. "
         "No se digitan códigos DANE.", False)])
vineta([("Las sedes que la visita técnica clasificó sin afectación vienen pre-marcadas y se "
         "confirman en bloque. ", False),
        ("Riosucio resuelve 30 de sus 92 sedes en un clic.", True)])
vineta([("Para las sedes afectadas se registran actividades, cantidades y valores unitarios. ",
         False),
        ("Todos los cálculos los hace el sistema", True),
        (": costo directo, administración, utilidad e IVA sobre la utilidad.", False)])
vineta([("La Secretaría verifica en pantalla propia, con el valor estimado al lado y la desviación "
         "calculada. Genera el formato oficial en PDF con firma.", False)])
vineta([("El municipio no ve la cifra estimada por la Secretaría: su presupuesto es una ", False),
        ("medición independiente", True),
        (", lo que permite validar el modelo donde ambas coincidan.", False)])

# ---------------------------------------------------------------- 3
titulo("3. Estado")
tabla([
    ["Componente", "Estado"],
    ["Aplicativo de las alcaldías", "Funcionando y probado"],
    ["Consola de verificación de la Secretaría", "Funcionando y probada"],
    ["Formato oficial en PDF", "Funcionando"],
    ["Catálogo de 975 sedes, 161 I.E. y 26 alcaldías", "Verificado contra la base maestra"],
    ["Enlaces con token para las 26 alcaldías", "Generados"],
    ["Instructivo para las alcaldías", "Listo"],
    ["Sitio de SharePoint y carpetas por alcaldía", "PENDIENTE — requiere autorización"],
    ["Automatización hacia el tablero de Power BI", "Especificada, falta construirla"],
], [10.0, 6.4])

# ---------------------------------------------------------------- 4
titulo("4. Verificación de calidad")
mixto([("El instrumento fue sometido a revisión antes de su uso. Se detectaron y corrigieron ",
        False), ("tres defectos", True), (":", False)], sa=3)
vineta([("Los campos rechazaban cifras escritas como ", False), ("1.500.000", True),
        (" y calculaban cero sin advertencia.", False)])
vineta([("En municipios grandes se perdían borradores sin que el usuario lo notara.", False)])
vineta([("Las fotografías se dañaban al reabrir una sede ya diligenciada.", False)])
mixto([("Los tres están corregidos y verificados. ", False),
       ("Se informan por transparencia: haberlos encontrado en pruebas y no en manos de un "
        "alcalde es parte del control de calidad.", True)], sa=4)

# ---------------------------------------------------------------- 5
titulo("5. Costo")
mixto([("Ninguno. ", True),
       ("No requiere comprar licencias ni contratar servicios: se utiliza la infraestructura "
        "Microsoft 365 con la que ya cuenta la Gobernación. La licencia Power Automate Premium "
        "fue descartada por innecesaria.", False)])

# ---------------------------------------------------------------- 6
titulo("6. Lo que se solicita")
tabla([
    ["Decisión", "Responsable"],
    ["Aval del instrumento para su uso con los municipios", "Dirección de Planeación"],
    ["Autorización del sitio de SharePoint y las carpetas por alcaldía", "Dirección de Planeación / TI"],
    ["Aval de los umbrales de alerta: Administración 12 %, Utilidad 8 %", "Dirección de Planeación"],
    ["Definición de los municipios para la prueba inicial", "Dirección de Planeación"],
], [11.2, 5.2])

p()
mixto([("Recomendación: ", True),
       ("iniciar con uno o dos municipios pequeños —Marulanda, 11 sedes, o San José, 13— con "
        "acompañamiento telefónico, antes del envío a las 26 alcaldías. Lo que se aprende de un "
        "alcalde diligenciando en vivo no se obtiene de más días de desarrollo.", False)], sa=10)

# ---------------------------------------------------------------- firmas
t = doc.add_table(rows=1, cols=2)
t.alignment = WD_TABLE_ALIGNMENT.LEFT
for j, (rot, sub) in enumerate([("Elaboró", "Práctica TIC — Secretaría de Educación de Caldas"),
                                ("Avala", "Dirección de Planeación")]):
    c = t.cell(0, j)
    c.text = ""
    par = c.paragraphs[0]
    par.add_run("\n\n_______________________________\n").font.size = Pt(9)
    r = par.add_run(rot + "\n"); r.bold = True; r.font.size = Pt(9)
    r2 = par.add_run(sub); r2.font.size = Pt(8); r2.font.color.rgb = GRIS
    t.columns[j].width = Cm(8.2)

rutas.asegurar_directorios()
doc.save(SALIDA)
print("Escrito:", SALIDA)
print("Párrafos:", len(doc.paragraphs), "| Tablas:", len(doc.tables))
